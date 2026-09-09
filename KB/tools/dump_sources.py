import io, sys, json
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
import openpyxl

p = r"C:\Users\Waseem\OneDrive\Desktop\SIH 2026\KB\raw\SIH26092_Knowledge_Base_Source_Register.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
print("SHEETS:", wb.sheetnames)
for ws in wb.worksheets:
    print("=" * 80)
    print("SHEET:", ws.title, "dims:", ws.dimensions, "max_row:", ws.max_row, "max_col:", ws.max_column)
    for row in ws.iter_rows(values_only=True):
        cells = [("" if c is None else str(c)) for c in row]
        print(" | ".join(cells))